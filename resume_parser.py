"""Module 1: Resume upload parsing - extract clean text from PDF / DOCX."""
import os
import re

import pdfplumber
from docx import Document

ALLOWED_EXTENSIONS = {"pdf", "docx"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def _read_pdf(path):
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _read_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    # many resumes keep content inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def clean_text(text):
    text = text.replace("\u2022", "\n- ").replace("\uf0b7", "\n- ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)
    return text.strip()


def extract_text(path):
    """Return cleaned text. Raises ValueError with a friendly message on failure."""
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")
    try:
        raw = _read_pdf(path) if ext == "pdf" else _read_docx(path)
    except Exception:
        raise ValueError("Could not read this file. It may be corrupted or password protected.")
    text = clean_text(raw)
    if len(text) < 50:
        raise ValueError("No readable text found. Scanned/image resumes are not supported.")
    return text
